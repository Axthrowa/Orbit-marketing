from __future__ import annotations

import csv
import json
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from etkilesim import cli, maintenance, notify, posts, providers, search, settings, system
from etkilesim.errors import PermissionDenied, ValidationFailed
from etkilesim.models import CollectResult, Metrics
from etkilesim.sentiment import LexiconEngine

from .helpers import make_user, public_resolver

ENGINE = LexiconEngine()


def add(conn: sqlite3.Connection, uid: int, url: str, t: datetime, **kw: object) -> int:
    return posts.add_post(conn, uid, url, now=t, resolver=public_resolver, **kw).post_id  # type: ignore[arg-type]


# ----- işler, sağlayıcılar, denetim -----------------------------------------------------------------


def test_record_job_success_and_failure(conn: sqlite3.Connection, t0: datetime) -> None:
    system.record_job(conn, "flush", "manual", lambda: "tamam", now=lambda: t0)
    with pytest.raises(RuntimeError):
        system.record_job(
            conn, "refresh", "manual", lambda: (_ for _ in ()).throw(RuntimeError("çöktü")), now=lambda: t0
        )
    with pytest.raises(ValidationFailed):
        system.record_job(conn, "nope", "manual", lambda: "x")
    jobs = system.jobs(conn)
    assert jobs["summary"] == {"ok": 1, "error": 1}
    assert jobs["latest"]["refresh"]["status"] == "error" and "çöktü" in jobs["latest"]["refresh"]["detail"]
    assert jobs["latest"]["flush"]["trigger"] == "manual"
    assert jobs["latest"]["discovery"]["status"] is None


def test_provider_health_circuit_states_and_success_rate(conn: sqlite3.Connection, t0: datetime) -> None:
    uid = make_user(conn, now=t0)
    ok = add(conn, uid, "https://x.com/a/status/1", t0)
    bad = add(conn, uid, "https://x.com/a/status/2", t0)
    add(conn, uid, "https://www.reddit.com/r/x/comments/abc/t/", t0)
    posts.record_collect(conn, ok, CollectResult("OK", Metrics(likes=1)), ENGINE, t0)
    posts.record_collect(conn, bad, CollectResult("DATA_UNAVAILABLE", blocked=True), ENGINE, t0)
    providers.record_failure(conn, "X", "bot", True, t0)
    health = {p["platform"]: p for p in system.provider_health(conn, t0)}
    assert health["X"]["circuit"] == "OPEN" and health["X"]["success_rate"] == 0.5
    assert health["Reddit"]["queue"] == 1 and health["Reddit"]["circuit"] == "CLOSED"
    later = {p["platform"]: p for p in system.provider_health(conn, t0 + timedelta(hours=1))}
    assert later["X"]["circuit"] == "HALF_OPEN"  # duraklatma bitti, sayaç henüz sıfırlanmadı


def test_retry_failed_requeues(conn: sqlite3.Connection, t0: datetime) -> None:
    uid = make_user(conn, now=t0)
    pid = add(conn, uid, "https://x.com/a/status/1", t0)
    posts.record_collect(conn, pid, CollectResult("DATA_UNAVAILABLE"), ENGINE, t0)
    notify.queue(conn, "strike", "1", "m", "k", uid, t0)
    conn.execute("UPDATE notifications SET status='failed', attempts=5")
    assert system.retry_failed(conn, "admin", t0) == {"notifications": 1, "posts": 1}
    assert conn.execute("SELECT status, attempts FROM notifications").fetchone()[:] == ("pending", 0)
    assert [r["id"] for r in posts.due_posts(conn, t0)] == [pid]
    assert system.audit_log(conn, action="system.")[0]["detail"] == {"notifications": 1, "posts": 1}


def test_system_commands_need_admin(conn: sqlite3.Connection, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv(cli.PIN_ENV, raising=False)
    for name in ("system-overview", "run-job", "retry-failed", "retention", "backup", "export"):
        with pytest.raises(PermissionDenied):
            cli.dispatch(name, {}, conn)


def test_manual_cli_jobs_are_recorded(conn: sqlite3.Connection, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv(cli.PIN_ENV, "246810")
    cli.dispatch("set-pin", {"new_pin": "246810"}, conn)
    cli.dispatch("run-job", {"name": "strike_check"}, conn)
    cli.dispatch("run-job", {"name": "retention"}, conn)
    rows = conn.execute("SELECT name, trigger, status FROM job_runs ORDER BY id").fetchall()
    assert [tuple(r) for r in rows] == [("strike_check", "manual", "ok"), ("retention", "manual", "ok")]
    with pytest.raises(ValidationFailed):
        cli.dispatch("run-job", {"name": "rm -rf"}, conn)


# ----- saklama, dışa aktarım, yedek ----------------------------------------------------------------


def test_retention_redacts_old_comment_text_but_keeps_sentiment(
    conn: sqlite3.Connection, t0: datetime
) -> None:
    uid = make_user(conn, now=t0)
    pid = add(conn, uid, "https://x.com/a/status/1", t0)
    posts.add_manual_comments(conn, pid, ["eski harika yorum"], ENGINE, "u", t0 - timedelta(days=200))
    posts.add_manual_comments(conn, pid, ["yeni berbat yorum"], ENGINE, "u", t0)
    conn.execute("INSERT INTO audit_log(actor,action,at) VALUES ('x','old','2025-01-01T00:00:00Z')")
    result = maintenance.apply_retention(conn, t0)
    assert result["comment_texts"] == 1 and result["audit"] >= 1
    texts = {r["text"]: r["label"] for r in conn.execute("SELECT text, label FROM comments")}
    assert texts[maintenance.REDACTED_TEXT] == "POSITIVE" and "yeni berbat yorum" in texts
    assert conn.execute("SELECT sentiment_sample FROM posts").fetchone()[0] == 2  # oranlar korunur
    settings.set_value(conn, "comment_retention_days", 0, "admin")
    assert "comment_texts" not in maintenance.apply_retention(conn, t0 + timedelta(days=999))


def test_export_formats_and_formula_injection(conn: sqlite3.Connection, t0: datetime, tmp_path: Path) -> None:
    uid = make_user(conn, "=HYPERLINK(evil)", now=t0)
    add(conn, uid, "https://x.com/a/status/1", t0, note="+cmd|calc")
    csv_out = maintenance.export(conn, "csv", tmp_path, t0)
    with open(csv_out["path"], encoding="utf-8-sig") as handle:
        row = next(csv.DictReader(handle, delimiter=";"))
    assert row["kisi"] == "'=HYPERLINK(evil)" and row["sayfa_notu"] == "'+cmd|calc"
    xlsx_out = maintenance.export(conn, "xlsx", tmp_path, t0)
    book = load_workbook(xlsx_out["path"])
    assert book.sheetnames == ["paylasimlar", "liderlik_30_gun", "topluluklar"]
    assert book.properties.creator == "created by axthrowa"
    json_out = maintenance.export(conn, "json", tmp_path, t0)
    data = json.loads(Path(json_out["path"]).read_text(encoding="utf-8"))
    assert data["meta"]["created_by"] == "axthrowa" and len(data["paylasimlar"]) == 1
    with pytest.raises(ValidationFailed):
        maintenance.export(conn, "pdf", tmp_path, t0)


def test_backup_is_a_valid_database_and_rotates(
    conn: sqlite3.Connection, t0: datetime, tmp_path: Path
) -> None:
    make_user(conn, now=t0)
    settings.set_value(conn, "backup_keep", 2, "admin")
    paths = [maintenance.backup(conn, tmp_path, t0 + timedelta(seconds=i))["path"] for i in range(3)]
    assert sorted(p.name for p in tmp_path.glob("etkilesim-*.db")) == sorted(Path(p).name for p in paths[1:])
    copy = sqlite3.connect(paths[-1])
    assert copy.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 1
    copy.close()


# ----- arama ----------------------------------------------------------------------------------------


def test_search_groups_and_scopes(conn: sqlite3.Connection, t0: datetime) -> None:
    a = make_user(conn, "Ayşe", now=t0)
    b = make_user(conn, "Burak", now=t0)
    add(conn, a, "https://x.com/a/status/1", t0, community_name="r/IndieGaming", note="indie notu")
    add(conn, b, "https://x.com/b/status/2", t0, community_name="r/IndieDev")
    everything = search.search(conn, "indie")
    assert len(everything["posts"]) == 2 and {c["name"] for c in everything["communities"]} == {
        "r/IndieGaming",
        "r/IndieDev",
    }
    mine = search.search(conn, "indie", user_id=a)
    assert [p["user_name"] for p in mine["posts"]] == ["Ayşe"] and mine["people"] == []
    assert search.search(conn, "%_")["posts"] == []  # joker karakterler kaçışlı
    with pytest.raises(ValidationFailed):
        search.search(conn, "a")
